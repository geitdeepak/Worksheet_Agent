"""Ready-to-fill Excel templates for every upload, so admins never have to guess the format."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = PatternFill("solid", fgColor="1B365D")


def build_xlsx(sheet: str, headers: list[str], rows: list[list], instructions: list[str], widths: list[int],
               text_cols: tuple[int, ...] = (), example_rows: int = 0, lists: dict[int, list[str]] | None = None,
               date_cols: tuple[int, ...] = ()) -> bytes:
    """rows: pre-filled rows; the first `example_rows` of them are greyed-out samples (ignored on import)."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = NAVY
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 22
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    for r in range(2, 2 + example_rows):
        for i in range(1, len(headers) + 1):
            ws.cell(row=r, column=i).font = Font(italic=True, color="8596AD")
    for r in range(2, 1002):
        for col in text_cols:  # keep IDs and phone numbers as text so Excel doesn't mangle them
            ws.cell(row=r, column=col).number_format = "@"
        for col in date_cols:
            ws.cell(row=r, column=col).number_format = "DD-MMM-YYYY"
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).alignment = Alignment(vertical="top", wrap_text=True)
    for col, options in (lists or {}).items():
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        letter = chr(64 + col)
        dv.add(f"{letter}2:{letter}1001")
    ws.freeze_panes = "A2"

    info = wb.create_sheet("How to fill this")
    for line in instructions:
        info.append([line])
    info.column_dimensions["A"].width = 110
    info["A1"].font = Font(bold=True, size=13, color="1B365D")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def datesheet_template(grade: str, sections: list[str], subjects: list[str]) -> bytes:
    sec = ", ".join(sections)
    rows = [["EXAMPLE-1", grade, sec, "Mathematics (example row, ignored)", "10-Oct-2026", "09:00"]]
    rows += [["", grade, sec, s, "", ""] for s in subjects]
    return build_xlsx(
        "Date sheet", ["Exam ID", "Class", "Section", "Subject", "Exam Date", "Exam Time"], rows,
        ["How to fill the date sheet", "",
         "• One row per exam.",
         "• Subject: use the same name as in Study material (e.g. Mathematics).",
         "• Exam Date: a date like 10-Oct-2026. Worksheets are sent two days before this date.",
         "• Section: leave as is for all sections, or write one section (A) if the exam is only for that section.",
         "• Exam Time and Exam ID are optional. An Exam ID is created for you if left blank.",
         "• The grey EXAMPLE row is ignored — you can leave it or delete it.",
         "", "After uploading you will see a preview and confirm it. Nothing changes until you confirm."],
        [12, 8, 10, 34, 14, 11], text_cols=(1,), example_rows=1, date_cols=(5,))


def syllabus_template(subjects: list[str]) -> bytes:
    rows = [["Science (example row, ignored)", "Chapter 1 Motion, Chapter 2 Force and Laws of Motion"]]
    rows += [[s, ""] for s in subjects]
    return build_xlsx(
        "Syllabus", ["Subject", "Syllabus (chapters and topics for this exam)"], rows,
        ["How to fill the exam syllabus", "",
         "• One row per subject. Your subjects are already listed — just fill in column B.",
         "• Write the chapters for this exam, e.g. 'Chapter 1, 2 and 5' or 'Ch 1-4'.",
         "• You can add topic names too, e.g. 'Chapter 3 Polynomials – remainder theorem only'.",
         "• To leave something out, write it on its own line: 'Chapter 6 is not included'.",
         "• Worksheets will only use these chapters. Questions outside them are blocked.",
         "• The grey example row is ignored.",
         "", "You can also upload the syllabus notice your school already has (PDF or Word) instead of this sheet."],
        [30, 90], example_rows=1)
