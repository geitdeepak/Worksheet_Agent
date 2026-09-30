"""Parse uploaded date sheets and student lists (.xlsx or .csv) into validated rows.

Parsing is deterministic on purpose: a parsing mistake here would trigger a worksheet on the
wrong day, so every row is either accepted as-is or returned as an issue for the admin to fix."""
import csv
import io
import re
from datetime import date, datetime

from openpyxl import load_workbook

from .common import normalize_whatsapp, valid_email

HEADER_ALIASES = {
    "class": {"class", "grade", "std", "standard"},
    "section": {"section", "sec", "sections"},
    "subject": {"subject", "paper", "subject name"},
    "exam_date": {"exam date", "date", "examdate", "exam_date"},
    "exam_time": {"exam time", "time", "examtime", "exam_time", "start time"},
    "exam_code": {"exam id", "examid", "exam code", "exam_id", "id"},
    "student_code": {"student id", "studentid", "student code", "student_id", "roll no", "roll number", "admission no"},
    "name": {"name", "student name", "full name"},
    "email": {"email", "email id", "e-mail", "email address", "mail"},
    "whatsapp": {"whatsapp", "whatsapp number", "whatsapp no", "phone", "mobile", "mobile number"},
    "active": {"active", "status"},
    "syllabus": {"syllabus", "topics", "chapters", "portion", "portions", "content", "syllabus topics",
                 "chapters topics", "chapters/topics", "syllabus/chapters", "scope",
                 "syllabus (chapters and topics for this exam)"},
}


def _norm_header(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower().replace("_", " "))


def _map_headers(headers: list) -> dict[int, str]:
    mapping = {}
    for i, h in enumerate(headers):
        nh = _norm_header(h)
        for field, aliases in HEADER_ALIASES.items():
            if nh in aliases or nh.replace(" ", "") in {a.replace(" ", "") for a in aliases}:
                mapping.setdefault(i, field)
                break
    return mapping


def read_table(filename: str, data: bytes) -> list[dict]:
    """Returns rows as {field: value, '_row': spreadsheet row number}."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        ws = wb.worksheets[0]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        raw = [r for r in csv.reader(io.StringIO(text))]
    else:
        raise ValueError("Upload an .xlsx or .csv file.")

    # The header is the first row that maps at least two known columns (sheets often have a title row).
    header_idx, mapping = None, {}
    for i, r in enumerate(raw[:10]):
        m = _map_headers(r)
        if len(m) >= 2:
            header_idx, mapping = i, m
            break
    if header_idx is None:
        raise ValueError("No header row found. Include columns such as Class, Subject and Exam date.")

    rows = []
    for offset, r in enumerate(raw[header_idx + 1:], start=header_idx + 2):
        if not any(c not in (None, "") and str(c).strip() for c in r):
            continue
        row = {"_row": offset}
        for i, field in mapping.items():
            row[field] = r[i] if i < len(r) else None
        rows.append(row)
    return rows


_DATE_FORMATS = ["%d-%b-%Y", "%d %b %Y", "%d-%B-%Y", "%d %B %Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y",
                 "%d.%m.%Y", "%d/%m/%y", "%d-%b-%y", "%b %d, %Y", "%B %d, %Y"]


def parse_date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", str(v).strip())
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_time(v) -> str | None:
    if v is None or v == "":
        return None
    if hasattr(v, "strftime"):
        return v.strftime("%H:%M")
    s = str(v).strip().upper()
    for fmt in ("%I:%M %p", "%I:%M%p", "%H:%M", "%H:%M:%S", "%I %p"):
        try:
            return datetime.strptime(s, fmt).strftime("%H:%M")
        except ValueError:
            continue
    return s


def _grade_of(v) -> str:
    """'Class 9', '9', 'IX', 9.0 → '9'."""
    s = str(v or "").strip().upper().replace("CLASS", "").replace("GRADE", "").replace("STD", "").strip(" .-")
    if s.endswith(".0"):
        s = s[:-2]
    roman = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7, "VIII": 8, "IX": 9, "X": 10, "XI": 11, "XII": 12}
    return str(roman.get(s, s))


def _sections(v) -> list[str]:
    return [p.strip().upper() for p in re.split(r"[,/&;]| and ", str(v or "")) if p.strip()]


def _subject_code(subject: str) -> str:
    words = re.findall(r"[A-Za-z]+", subject)
    if len(words) >= 2:
        return "".join(w[0] for w in words).upper()[:4]
    return (words[0][:4] if words else "SUBJ").upper()


def parse_datesheet(filename: str, data: bytes, grade: str, class_sections: list[str]) -> tuple[list[dict], list[dict]]:
    """Returns (exams, issues). An issue is {row, message}; rows with issues are not imported."""
    rows = read_table(filename, data)
    exams, issues, seen = [], [], set()
    counters: dict[str, int] = {}
    for r in rows:
        n = r["_row"]
        if str(r.get("exam_code") or "").strip().upper().startswith("EXAMPLE"):
            continue  # template sample row
        problems = []
        subject = str(r.get("subject") or "").strip()
        if not subject:
            problems.append("subject is missing")
        cls = r.get("class")
        if cls not in (None, "") and _grade_of(cls) != str(grade):
            problems.append(f"belongs to Class {_grade_of(cls)}, not Class {grade}")
        d = parse_date(r.get("exam_date"))
        if d is None and r.get("exam_date") in (None, ""):
            problems.append("no exam date, so this row was skipped")
        elif d is None:
            problems.append(f"exam date '{r.get('exam_date')}' is not a date (write it like 10-Oct-2026)")
        secs = _sections(r.get("section"))
        unknown = [s for s in secs if class_sections and s not in class_sections]
        if unknown:
            problems.append(f"section {', '.join(unknown)} is not a section of this class")
        if problems:
            issues.append({"row": n, "message": f"Row {n} · " + "; ".join(problems)})
            continue
        code = str(r.get("exam_code") or "").strip()
        if not code:
            sc = _subject_code(subject)
            counters[sc] = counters.get(sc, 0) + 1
            code = f"EX{grade}-{sc}-{counters[sc]:03d}"
        key = (subject.lower(), d, ",".join(secs))
        if key in seen:
            issues.append({"row": n, "message": f"Row {n} · duplicates an earlier row ({subject}, {d.isoformat()})"})
            continue
        seen.add(key)
        exams.append({"exam_code": code, "subject_name": subject, "sections": ", ".join(secs),
                      "exam_date": d, "exam_time": parse_time(r.get("exam_time"))})
    exams.sort(key=lambda e: (e["exam_date"], e["subject_name"]))
    return exams, issues


def parse_students(filename: str, data: bytes, grade: str, class_sections: list[str],
                   default_cc: str = "91") -> tuple[list[dict], list[dict]]:
    return parse_student_rows(read_table(filename, data), grade, class_sections, default_cc)


def grade_of(v) -> str:
    return _grade_of(v)


def parse_student_rows(rows: list[dict], grade: str, class_sections: list[str],
                       default_cc: str = "91") -> tuple[list[dict], list[dict]]:
    """Validate student rows for one class. Returns (students, issues); rows with errors are skipped,
    rows with only warnings (bad email/number) are imported and reported."""
    students, issues, seen = [], [], set()
    for r in rows:
        n = r["_row"]
        code = str(r.get("student_code") or "").strip()
        if code.endswith(".0"):
            code = code[:-2]
        name = str(r.get("name") or "").strip()
        if code.upper() == "EXAMPLE" or code.upper().startswith("EXAMPLE-"):
            continue  # template sample rows
        problems = []
        if not code:
            problems.append("student ID is missing")
        if not name:
            problems.append("name is missing")
        cls = r.get("class")
        if cls not in (None, "") and _grade_of(cls) != str(grade):
            problems.append(f"belongs to Class {_grade_of(cls)}, not Class {grade}")
        secs = _sections(r.get("section"))
        section = secs[0] if secs else (class_sections[0] if len(class_sections) == 1 else "")
        if not section:
            problems.append("section is missing")
        elif class_sections and section not in class_sections:
            problems.append(f"section {section} is not a section of this class")
        if code in seen:
            problems.append(f"student ID {code} appears twice")
        if problems:
            issues.append({"row": n, "message": f"Row {n} · " + "; ".join(problems)})
            continue
        seen.add(code)
        email = str(r.get("email") or "").strip() or None
        wa_raw = r.get("whatsapp")
        whatsapp = normalize_whatsapp(str(wa_raw) if wa_raw not in (None, "") else None, default_cc)
        warn = []
        if email and not valid_email(email):
            warn.append("email looks invalid; delivery to it will fail")
        if wa_raw not in (None, "") and whatsapp is None:
            warn.append("WhatsApp number is invalid; WhatsApp will be skipped")
        active_raw = str(r.get("active") if r.get("active") is not None else "yes").strip().lower()
        students.append({"student_code": code, "name": name, "section": section, "email": email,
                         "whatsapp": whatsapp, "active": active_raw not in {"no", "false", "0", "inactive", "left"}})
        if warn:
            issues.append({"row": n, "message": f"Row {n} · {code} · " + "; ".join(warn), "warning": True})
    return students, issues
